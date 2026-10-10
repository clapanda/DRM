from __future__ import annotations

import argparse
import struct
import zlib
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def create_fixture(output: str | Path) -> Path:
    destination = Path(output)
    destination.parent.mkdir(parents=True, exist_ok=True)
    image_path = destination.with_suffix(".fixture-image.png")
    _write_test_png(image_path)

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    for style_name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3"]:
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    document.styles["Normal"].font.size = Pt(10.5)

    title = document.add_paragraph(style="Title")
    title.add_run("DRM 人工测试需求文档")
    intro = document.add_paragraph(
        "本文档仅包含合成测试内容，用于验证解析、定位、审批修改和自动验证流程。"
    )
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    document.add_heading("现金流指标", level=1)
    document.add_heading("余额类指标", level=2)
    document.add_paragraph("FE250 重定价前余额：本句是待完善的合成占位规则。")
    document.add_paragraph("相邻规则：余额类指标应保留来源字段和计算日期。")
    document.add_paragraph("重复规则")

    table = document.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    headers = ["指标码", "名称", "测试说明"]
    for column, value in enumerate(headers):
        table.cell(0, column).text = value
    table.cell(1, 0).text = "FE250"
    table.cell(1, 1).text = "重定价前余额"
    table.cell(1, 2).text = "仅用于定位测试，不代表真实业务口径"
    table.cell(2, 0).text = "FE251"
    table.cell(2, 1).text = "示例指标"
    table.cell(2, 2).text = "明细指标展开规则示例"
    merged = table.cell(3, 0).merge(table.cell(3, 1))
    merged.text = "合并单元格"
    table.cell(3, 2).text = "用于结构解析测试"

    document.add_heading("重定价规则", level=1)
    document.add_heading("事件处理", level=2)
    document.add_paragraph("重定价现金流事件的业务口径需要以正式来源为依据。")
    document.add_paragraph("重复规则")
    picture_paragraph = document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.add_run().add_picture(str(image_path), width=Cm(10))
    caption = document.add_paragraph("图 1 合成测试图片")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("附录", level=1)
    document.add_heading("连续版本验证", level=2)
    document.add_paragraph("版本一基线内容。")
    document.save(destination)
    image_path.unlink()
    return destination


def _write_test_png(path: Path, width: int = 640, height: int = 180) -> None:
    def chunk(kind: bytes, payload: bytes) -> bytes:
        return (
            struct.pack(">I", len(payload))
            + kind
            + payload
            + struct.pack(">I", zlib.crc32(kind + payload) & 0xFFFFFFFF)
        )

    rows = []
    for y in range(height):
        row = bytearray([0])
        for x in range(width):
            border = x < 8 or x >= width - 8 or y < 8 or y >= height - 8
            stripe = 230 < x < 410 and 70 < y < 110
            color = (79, 129, 189) if border or stripe else (220, 230, 241)
            row.extend(color)
        rows.append(bytes(row))
    payload = b"".join(rows)
    png = (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(payload, level=9))
        + chunk(b"IEND", b"")
    )
    path.write_bytes(png)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("output")
    args = parser.parse_args()
    path = create_fixture(args.output)
    print(path.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
