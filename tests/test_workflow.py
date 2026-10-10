from __future__ import annotations

import json
import shutil
import unittest
import uuid
from pathlib import Path

from docx import Document

from drm.approval import approve_changes
from drm.docx_parser import export_inspection, parse_docx
from drm.editor import apply_plan
from drm.errors import SafetyError
from drm.hashing import sha256_file
from drm.io_utils import write_json
from drm.models import ChangePlan
from drm.planner import build_plan
from drm.search import search_document
from drm.validator import validate_document
from tools.create_test_fixture import create_fixture


class WorkflowTests(unittest.TestCase):
    def setUp(self) -> None:
        local_temp = Path(__file__).parent / ".tmp"
        local_temp.mkdir(exist_ok=True)
        self.root = local_temp / f"case-{uuid.uuid4().hex}"
        self.root.mkdir()
        self.source = create_fixture(self.root / "source.docx")
        self.original_hash = sha256_file(self.source)

    def tearDown(self) -> None:
        shutil.rmtree(self.root, ignore_errors=True)

    def test_inspect_exports_structure_media_and_profile_hits(self) -> None:
        output = self.root / "inspection"
        paths = export_inspection(self.source, output, profile="ofs-alm")
        self.assertTrue(all(Path(path).exists() for path in paths.values()))
        structure = json.loads((output / "structure.json").read_text(encoding="utf-8"))
        media = json.loads((output / "media_index.json").read_text(encoding="utf-8"))
        codes = {item["code"] for item in structure["profile_analysis"]["indicators"]}
        self.assertIn("FE250", codes)
        self.assertGreaterEqual(structure["document"]["table_count"], 1)
        table = next(block for block in structure["blocks"] if block["kind"] == "table")
        self.assertEqual(table["rows"][3][0]["grid_span"], 2)
        self.assertEqual(len(media["media"]), 1)
        self.assertTrue(Path(media["media"][0]["extracted_path"]).exists())
        self.assertEqual(sha256_file(self.source), self.original_hash)

    def test_search_finds_paragraph_and_table_cell(self) -> None:
        result = search_document(self.source, "FE250", profile="ofs-alm")
        kinds = {match["kind"] for match in result["matches"]}
        self.assertEqual(kinds, {"paragraph", "table"})
        self.assertGreaterEqual(result["match_count"], 2)

    def test_plan_marks_ambiguous_target_as_conflict(self) -> None:
        proposal = self.root / "ambiguous.json"
        write_json(
            proposal,
            {
                "plan_id": "ambiguous-plan",
                "changes": [
                    self._change(
                        "C-AMB",
                        match_text="重复规则",
                        operation="replace_paragraph",
                        content="不得执行的替换",
                    )
                ],
            },
        )
        plan = build_plan(
            self.source, proposal, self.root / "ambiguous-plan.json", "v1", profile="ofs-alm"
        )
        self.assertEqual(plan.changes[0].planning_status.value, "conflict")
        self.assertEqual(plan.changes[0].approval_status.value, "pending")
        self.assertEqual(sha256_file(self.source), self.original_hash)

    def test_plan_resets_proposal_approval_until_explicit_approval_step(self) -> None:
        structure = parse_docx(self.source)
        target = next(
            block
            for block in structure["blocks"]
            if block["kind"] == "paragraph" and block["text"].startswith("相邻规则")
        )
        change = self._change(
            "C-RESET",
            anchor_id=target["anchor_id"],
            operation="insert_after",
            content="必须重新审批。",
        )
        change["approval_status"] = "approved"
        proposal = self.root / "preapproved-proposal.json"
        write_json(proposal, {"changes": [change]})
        plan = build_plan(self.source, proposal, self.root / "reset-plan.json", "v1")
        self.assertEqual(plan.changes[0].approval_status.value, "pending")
        self.assertTrue(any("reset" in note for note in plan.changes[0].planning_notes))

    def test_approved_apply_and_validate_preserve_unrelated_content_and_media(self) -> None:
        plan_path = self._build_approved_plan(self.source, self.root / "plan.json")
        output = self.root / "output.docx"
        audit = apply_plan(self.source, plan_path, output)
        report = validate_document(
            self.source, output, plan_path, self.root / "validation.json"
        )
        self.assertEqual(audit["status"], "partial")
        self.assertEqual(len(audit["applied_changes"]), 2)
        self.assertTrue(any(item["reason"] == "not_approved" for item in audit["skipped_changes"]))
        self.assertEqual(report["status"], "valid")
        self.assertTrue(report["checks"]["media_integrity"]["preserved"])
        self.assertEqual(report["checks"]["unapproved_content_change_count"], 0)
        output_text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
        self.assertIn("经批准的合成规则", output_text)
        self.assertIn("经批准的相邻补充", output_text)
        self.assertNotIn("未批准内容", output_text)
        self.assertEqual(sha256_file(self.source), self.original_hash)

    def test_validator_detects_unapproved_change(self) -> None:
        plan_path = self._build_approved_plan(self.source, self.root / "plan.json")
        output = self.root / "output.docx"
        apply_plan(self.source, plan_path, output)
        tampered = Document(output)
        target = next(p for p in tampered.paragraphs if p.text == "版本一基线内容。")
        target.text = "未经批准的附录改动。"
        tampered.save(output)
        report = validate_document(
            self.source, output, plan_path, self.root / "tampered-validation.json"
        )
        self.assertEqual(report["status"], "invalid")
        self.assertGreater(report["checks"]["unapproved_content_change_count"], 0)

    def test_validator_detects_unapproved_format_change(self) -> None:
        plan_path = self._build_approved_plan(self.source, self.root / "plan.json")
        output = self.root / "output.docx"
        apply_plan(self.source, plan_path, output)
        tampered = Document(output)
        target = next(p for p in tampered.paragraphs if p.text == "版本一基线内容。")
        target.runs[0].bold = True
        tampered.save(output)
        report = validate_document(
            self.source, output, plan_path, self.root / "format-validation.json"
        )
        self.assertEqual(report["status"], "invalid")
        self.assertGreater(report["checks"]["unapproved_content_change_count"], 0)

    def test_consecutive_versions_reanchor_against_new_hash(self) -> None:
        plan_one = self._build_approved_plan(self.source, self.root / "plan-one.json")
        version_one = self.root / "version-one.docx"
        apply_plan(self.source, plan_one, version_one)

        structure = parse_docx(version_one)
        target = next(
            block
            for block in structure["blocks"]
            if block["kind"] == "paragraph" and block["text"] == "版本一基线内容。"
        )
        proposal = self.root / "version-two-proposal.json"
        change = self._change(
            "C-V2",
            anchor_id=target["anchor_id"],
            operation="replace_paragraph",
            content="版本二经批准内容。",
        )
        write_json(proposal, {"plan_id": "version-two-plan", "changes": [change]})
        plan_two_path = self.root / "plan-two.json"
        draft_two_path = self.root / "plan-two-draft.json"
        plan_two = build_plan(version_one, proposal, draft_two_path, "v2")
        approve_changes(
            draft_two_path,
            plan_two_path,
            ["C-V2"],
            approver="automated-test",
        )
        version_two = self.root / "version-two.docx"
        apply_plan(version_one, plan_two_path, version_two)
        report = validate_document(
            version_one, version_two, plan_two_path, self.root / "version-two-validation.json"
        )
        self.assertEqual(report["status"], "valid")
        self.assertEqual(plan_two.target_document.sha256, sha256_file(version_one))

    def test_apply_rejects_manually_forged_approval_status(self) -> None:
        structure = parse_docx(self.source)
        target = next(
            block
            for block in structure["blocks"]
            if block["kind"] == "paragraph" and block["text"].startswith("相邻规则")
        )
        proposal = self.root / "forged-proposal.json"
        write_json(
            proposal,
            {
                "changes": [
                    self._change(
                        "C-FORGED",
                        anchor_id=target["anchor_id"],
                        operation="insert_after",
                        content="伪造审批不得执行。",
                    )
                ]
            },
        )
        plan_path = self.root / "forged-plan.json"
        build_plan(self.source, proposal, plan_path, "v1")
        payload = json.loads(plan_path.read_text(encoding="utf-8"))
        payload["changes"][0]["approval_status"] = "approved"
        write_json(plan_path, payload)
        output = self.root / "forged-output.docx"

        with self.assertRaises(SafetyError):
            apply_plan(self.source, plan_path, output)
        self.assertFalse(output.exists())
        self.assertEqual(sha256_file(self.source), self.original_hash)

    def test_apply_rejects_content_changed_after_approval(self) -> None:
        plan_path = self._build_approved_plan(self.source, self.root / "approved.json")
        payload = json.loads(plan_path.read_text(encoding="utf-8"))
        payload["changes"][0]["proposed_content"] = "审批后被篡改的内容。"
        write_json(plan_path, payload)
        output = self.root / "tampered-plan-output.docx"

        with self.assertRaises(SafetyError):
            apply_plan(self.source, plan_path, output)
        self.assertFalse(output.exists())
        self.assertEqual(sha256_file(self.source), self.original_hash)

    def _build_approved_plan(self, source: Path, output: Path) -> Path:
        structure = parse_docx(source)
        first = next(
            block
            for block in structure["blocks"]
            if block["kind"] == "paragraph" and block["text"].startswith("FE250 重定价前余额")
        )
        adjacent = next(
            block
            for block in structure["blocks"]
            if block["kind"] == "paragraph" and block["text"].startswith("相邻规则")
        )
        proposal = self.root / f"{output.stem}-proposal.json"
        changes = [
            self._change(
                "C-REPLACE",
                anchor_id=first["anchor_id"],
                operation="replace_paragraph",
                content="FE250 重定价前余额：经批准的合成规则，仅用于测试。",
            ),
            self._change(
                "C-INSERT",
                anchor_id=adjacent["anchor_id"],
                operation="insert_after",
                content="经批准的相邻补充。",
            ),
            self._change(
                "C-PENDING",
                anchor_id=adjacent["anchor_id"],
                operation="insert_before",
                content="未批准内容",
            ),
        ]
        write_json(proposal, {"plan_id": "approved-test-plan", "changes": changes})
        draft_path = output.with_name(output.stem + "-draft.json")
        build_plan(source, proposal, draft_path, "v1", profile="ofs-alm")
        approve_changes(
            draft_path,
            output,
            ["C-REPLACE", "C-INSERT"],
            approver="automated-test",
        )
        ChangePlan.load(output)
        return output

    @staticmethod
    def _change(
        change_id: str,
        *,
        operation: str,
        content: str,
        anchor_id: str | None = None,
        match_text: str | None = None,
    ) -> dict[str, object]:
        return {
            "change_id": change_id,
            "source_document": "synthetic-source.docx",
            "source_location": "synthetic evidence paragraph",
            "target_location": {
                "block_kind": "paragraph",
                "anchor_id": anchor_id,
                "match_text": match_text,
            },
            "operation": operation,
            "proposed_content": content,
            "rationale": "Synthetic automated test",
            "evidence": [
                {
                    "reference": "synthetic-source.docx#test",
                    "excerpt": "Synthetic evidence",
                    "evidence_type": "test_fixture",
                }
            ],
            "confidence_status": "confirmed_by_source",
            "approval_status": "pending",
            "validation_rules": [
                {"rule": "exact_text_present", "expected": content}
            ],
        }


if __name__ == "__main__":
    unittest.main()
